import docx
import openpyxl
import os
import sys
import argparse

def audit_docx(docx_path):
    print(f"\n[AUDIT DOCX] Memeriksa format & struktur naskah: {docx_path}")
    doc = docx.Document(docx_path)
    findings = []
    score = 100

    # 1. Check Margins
    sec = doc.sections[0]
    top_cm = round(sec.top_margin.cm, 2)
    bot_cm = round(sec.bottom_margin.cm, 2)
    left_cm = round(sec.left_margin.cm, 2)
    right_cm = round(sec.right_margin.cm, 2)
    if all(2.4 <= m <= 2.6 for m in [top_cm, bot_cm, left_cm, right_cm]):
        findings.append(("PASS", f"Margin halaman standar (Normal): Top={top_cm}cm, Bottom={bot_cm}cm, Left={left_cm}cm, Right={right_cm}cm"))
    else:
        findings.append(("WARN", f"Margin tidak persis 2.54 cm: Top={top_cm}cm, Bot={bot_cm}cm, Left={left_cm}cm, Right={right_cm}cm (-5 poin)"))
        score -= 5

    # 2. Check Font
    normal_font = doc.styles['Normal'].font.name
    if normal_font and "Times" in normal_font:
        findings.append(("PASS", f"Font naskah utama: {normal_font} (Sesuai FEB UKRIDA)"))
    else:
        findings.append(("WARN", f"Font utama bukan Times New Roman: {normal_font} (-5 poin)"))
        score -= 5

    # 3. Check Chapters
    full_text = "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
    chapters_required = [
        ("BAB 1 / BAB I", ["BAB 1", "BAB I"]),
        ("BAB II (Pemasaran)", ["BAB II", "ASPEK PEMASARAN"]),
        ("BAB III (Manajemen)", ["BAB III", "ASPEK MANAJEMEN"]),
        ("BAB IV (Teknis Operasi)", ["BAB IV", "ASPEK TEKNIS"]),
        ("BAB V (Keuangan)", ["BAB V", "ASPEK KEUANGAN", "APEK KEUANGAN"]),
        ("BAB VI (Kesimpulan)", ["BAB VI", "KESIMPULAN"])
    ]
    for ch_name, keywords in chapters_required:
        if any(kw in full_text for kw in keywords):
            findings.append(("PASS", f"Bab terdeteksi: {ch_name}"))
        else:
            findings.append(("FAIL", f"Bab hilang / tidak terdeteksi: {ch_name} (-10 poin)"))
            score -= 10

    # 4. Check Mandatory Tables & Figures
    mandatory_items = [
        ("Tabel 1.1 (Jenis Produk)", ["Tabel 1.1"]),
        ("Tabel 4.1 (Peralatan Capex)", ["Tabel 4.1"]),
        ("Tabel 4.2 (Gantt Chart)", ["Tabel 4.2"]),
        ("Tabel 5.1 (NCF & Outlay)", ["Tabel 5.1"]),
        ("Tabel 6.1 (Kesimpulan)", ["Tabel 6.1"]),
        ("Gambar 1.1 (Visual Produk)", ["Gambar 1.1"]),
        ("Gambar 2.1 (Lokasi)", ["Gambar 2.1"]),
        ("Gambar 3.1 (Struktur Organisasi)", ["Gambar 3.1"]),
        ("Gambar 4.1 (Layout Ruang)", ["Gambar 4.1"]),
        ("Gambar 4.2 (Network Planning)", ["Gambar 4.2"])
    ]
    for item_name, kw_list in mandatory_items:
        if any(kw in full_text for kw in kw_list):
            findings.append(("PASS", f"Elemen ditemukan: {item_name}"))
        else:
            findings.append(("WARN", f"Elemen tabel/gambar belum lengkap: {item_name} (-3 poin)"))
            score -= 3

    return score, findings

def audit_xlsx(xlsx_path):
    print(f"\n[AUDIT XLSX] Memeriksa model finansial: {xlsx_path}")
    wb = openpyxl.load_workbook(xlsx_path, data_only=False)
    wb_data = openpyxl.load_workbook(xlsx_path, data_only=True)
    findings = []
    score = 100

    # Check Sheet existence
    expected_sheets = ["Sheet1", "Sheet2", "Sheet3"]
    for s in expected_sheets:
        if s in wb.sheetnames:
            findings.append(("PASS", f"Sheet terverifikasi: {s}"))
        else:
            findings.append(("FAIL", f"Sheet wajib hilang: {s} (-15 poin)"))
            score -= 15

    if "Sheet3" in wb.sheetnames:
        ws3 = wb["Sheet3"]
        ws3_d = wb_data["Sheet3"]

        # Check Initial Outlay
        outlay_formula = str(ws3['C9'].value or '')
        if "SUM" in outlay_formula or "Sheet1" in outlay_formula:
            findings.append(("PASS", f"Formula Total Outlay terhubung: {outlay_formula}"))
        else:
            findings.append(("WARN", f"Total Outlay tidak menggunakan formula otomatis: {outlay_formula} (-5 poin)"))
            score -= 5

        # Check NPV, IRR, PP, PI
        npv_val = ws3_d['D34'].value
        irr_val = ws3_d['C44'].value
        pp_val = ws3_d['E49'].value
        pi_val = ws3_d['C55'].value

        findings.append(("INFO", f"Nilai Terhitung: NPV={npv_val}, IRR={irr_val}, PP={pp_val}, PI={pi_val}"))

        # Feasibility check
        if npv_val is not None:
            try:
                if float(npv_val) > 0:
                    findings.append(("PASS", f"Kriteria NPV: {float(npv_val):,.0f} > 0 (LAYAK)"))
                else:
                    findings.append(("FAIL", f"Kriteria NPV tidak layak: {npv_val} <= 0 (-15 poin)"))
                    score -= 15
            except Exception:
                pass

        if irr_val is not None:
            try:
                if float(irr_val) > 0.20:
                    findings.append(("PASS", f"Kriteria IRR: {float(irr_val)*100:.2f}% > 20.00% (LAYAK)"))
                else:
                    findings.append(("FAIL", f"Kriteria IRR tidak layak: {irr_val} <= 20% (-15 poin)"))
                    score -= 15
            except Exception:
                pass

        if pi_val is not None:
            try:
                if float(pi_val) > 1.20:
                    findings.append(("PASS", f"Kriteria PI: {float(pi_val):.2f} > 1.20 (LAYAK)"))
                else:
                    findings.append(("FAIL", f"Kriteria PI tidak layak: {pi_val} <= 1.20 (-10 poin)"))
                    score -= 10
            except Exception:
                pass

    return score, findings

def main():
    parser = argparse.ArgumentParser(description="Audit Dokumen Tugas SKB FEB UKRIDA")
    parser.add_argument("--docx", help="Path ke file Word .docx", default=None)
    parser.add_argument("--xlsx", help="Path ke file Excel .xlsx", default=None)
    args = parser.parse_args()

    # Fallback to defaults if no args
    if not args.docx and not args.xlsx:
        default_docx = r"d:\Perkuliahan\Kelass\SKB\Laporan_SKB_Template.docx"
        default_xlsx = r"d:\Perkuliahan\Kelass\SKB\Model_Finansial_Template.xlsx"
        if os.path.exists(default_docx):
            args.docx = default_docx
        if os.path.exists(default_xlsx):
            args.xlsx = default_xlsx

    print("=" * 65)
    print("      SKB COMPREHENSIVE QUALITY AUDITOR REPORT")
    print("      Standar: Fakultas Ekonomi & Bisnis UKRIDA")
    print("=" * 65)

    total_score = 0
    components = 0

    if args.docx and os.path.exists(args.docx):
        score_doc, findings_doc = audit_docx(args.docx)
        total_score += score_doc
        components += 1
        for status, msg in findings_doc:
            badge = f"[{status}]"
            print(f"{badge:<8} {msg}")

    if args.xlsx and os.path.exists(args.xlsx):
        score_xls, findings_xls = audit_xlsx(args.xlsx)
        total_score += score_xls
        components += 1
        for status, msg in findings_xls:
            badge = f"[{status}]"
            print(f"{badge:<8} {msg}")

    final_score = (total_score / components) if components > 0 else 0
    print("\n" + "=" * 65)
    print(f"SKOR AKHIR AUDIT MUTU: {final_score:.1f} / 100.0")
    if final_score >= 85:
        print("PREDIKAT: SANGAT BAIK (READY TO SUBMIT)")
    elif final_score >= 70:
        print("PREDIKAT: CUKUP (PERLU PERBAIKAN MINOR)")
    else:
        print("PREDIKAT: KURANG (REVISI MAYOR SEBELUM PENGUMPULAN)")
    print("=" * 65 + "\n")

if __name__ == "__main__":
    main()
