import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re
import os

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = f'w:{edge}'
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key in ["sz", "val", "color", "space", "shadow"]:
                if key in edge_data:
                    element.set(qn(f'w:{key}'), str(edge_data[key]))

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color_hex)
    tcPr.append(shd)

def fix_document():
    doc_path = r"d:\Perkuliahan\Kelass\SKB\skb debby MARTABAK HATI\MARTABAK HATI\Martabak Hati_ORIGINAL_BACKUP.docx"
    doc = docx.Document(doc_path)

    # 1. Ensure Normal Style is Times New Roman 12pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)

    # 2. Fix Cover Page (Paragraphs 0 to 19)
    # Clear existing cover paragraphs
    for i in range(20):
        doc.paragraphs[i].text = ""

    # Rebuild Cover Page
    p0 = doc.paragraphs[0]
    p0.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p0.paragraph_format.space_before = Pt(24)
    p0.paragraph_format.space_after = Pt(6)
    r0 = p0.add_run("Studi Kelayakan Bisnis")
    r0.font.name = "Times New Roman"
    r0.font.size = Pt(14)
    r0.bold = True

    p1 = doc.paragraphs[1]
    p1.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p1.paragraph_format.space_after = Pt(18)
    r1 = p1.add_run("“ MARTABAK HATI ”")
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(18)
    r1.bold = True

    # Paragraph 2: Insert Logo UKRIDA
    p2 = doc.paragraphs[2]
    p2.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p2.paragraph_format.space_after = Pt(24)
    logo_path = r"d:\Perkuliahan\Kelass\SKB\logo_ukrida.png"
    if os.path.exists(logo_path):
        run_logo = p2.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.0))

    # Paragraph 3: Disusun Oleh & Team Members
    p3 = doc.paragraphs[3]
    p3.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p3.paragraph_format.space_after = Pt(4)
    r_disusun = p3.add_run("Disusun Oleh:")
    r_disusun.font.name = "Times New Roman"
    r_disusun.font.size = Pt(12)
    r_disusun.bold = True

    members = [
        "Arthur Reezan (312023002)",
        "Jennese Putra Alamsyah Sukadi (312023033)",
        "Valendrik Dwiputra Wirawan (312023013)",
        "Affandy (312023075)",
        "Steven Putra Tjhin (312023015)"
    ]
    p4 = doc.paragraphs[4]
    p4.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p4.paragraph_format.space_after = Pt(40)
    for m in members:
        r_m = p4.add_run(f"{m}\n")
        r_m.font.name = "Times New Roman"
        r_m.font.size = Pt(11)

    # Paragraph 5: Institusi
    p5 = doc.paragraphs[5]
    p5.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p5.paragraph_format.space_before = Pt(30)
    p5.paragraph_format.space_after = Pt(4)
    r5 = p5.add_run("Program Studi Manajemen\nFakultas Ekonomi & Bisnis\nUniversitas Kristen Krida Wacana\n2023/2024")
    r5.font.name = "Times New Roman"
    r5.font.size = Pt(13)
    r5.bold = True

    # 3. Clean up and standardize all paragraphs from P20 onwards
    for idx in range(20, len(doc.paragraphs)):
        p = doc.paragraphs[idx]
        txt = p.text

        # Fix Typos in text
        if "APEK KEUANGAN" in txt:
            txt = txt.replace("APEK KEUANGAN", "ASPEK KEUANGAN")
        if "Cah Inflow" in txt:
            txt = txt.replace("Cah Inflow", "Cash Inflow")
        if "Cah Outflow" in txt:
            txt = txt.replace("Cah Outflow", "Cash Outflow")
        if "Intial Outlay" in txt:
            txt = txt.replace("Intial Outlay", "Initial Outlay")
        if "10 martabak dengan harga Rp 7.000" in txt:
            txt = txt.replace("10 martabak dengan harga Rp 7.000", "10 gelas dengan harga Rp 7.000")
        if "12 martabak dengan harga Rp 5.000" in txt:
            txt = txt.replace("12 martabak dengan harga Rp 5.000", "12 gelas dengan harga Rp 5.000")
        if "10 martabak dengan harga Rp 5.000 /pcs" in txt:
            txt = txt.replace("10 martabak dengan harga Rp 5.000 /pcs", "10 botol dengan harga Rp 5.000 /pcs")
        if "Tabel 5.25. Profitability Index" in txt:
            txt = txt.replace("Tabel 5.25. Profitability Index", "Tabel 5.26. Profitability Index")

        # Re-apply text if modified
        if txt != p.text:
            p.text = txt

        # Force font name to Times New Roman on all runs
        for r in p.runs:
            r.font.name = "Times New Roman"
            if not r.font.size:
                r.font.size = Pt(12)

    # 4. Standardize all 23 Tables (Convert Calibri -> Times New Roman, fix spacing and borders)
    border_kwargs = {
        'top': {'sz': 4, 'val': 'single', 'color': '808080'},
        'bottom': {'sz': 4, 'val': 'single', 'color': '808080'},
        'left': {'sz': 4, 'val': 'single', 'color': '808080'},
        'right': {'sz': 4, 'val': 'single', 'color': '808080'},
        'insideH': {'sz': 4, 'val': 'single', 'color': 'D0D0D0'},
        'insideV': {'sz': 4, 'val': 'single', 'color': 'D0D0D0'},
    }

    for t_idx, t in enumerate(doc.tables):
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for row_idx, row in enumerate(t.rows):
            is_header = (row_idx == 0)
            for c_idx, cell in enumerate(row.cells):
                set_cell_border(cell, **border_kwargs)
                if is_header:
                    set_cell_shading(cell, "D9E1F2")

                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.0
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.space_before = Pt(2)

                    # Clean multiple spaces in cell text (e.g. 'Rp          20.000' -> 'Rp 20.000')
                    raw_txt = p.text
                    clean_txt = re.sub(r'\s{2,}', ' ', raw_txt).strip()
                    if clean_txt != raw_txt:
                        p.text = clean_txt

                    # Check alignment: if starts with Rp or is number, align right
                    if clean_txt.startswith("Rp") or clean_txt.startswith("-Rp") or (clean_txt.replace(".", "").replace(",", "").replace("%", "").isdigit() and not is_header and len(clean_txt) > 2):
                        p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT
                    elif clean_txt.isdigit() and len(clean_txt) <= 2:
                        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

                    for r in p.runs:
                        r.font.name = "Times New Roman"
                        r.font.size = Pt(10)
                        if is_header:
                            r.bold = True

    # Save to both target file and a clean separate file
    target_path = r"d:\Perkuliahan\Kelass\SKB\skb debby MARTABAK HATI\MARTABAK HATI\Martabak Hati_312020064_DebbyDelicia.docx"
    clean_path = r"d:\Perkuliahan\Kelass\SKB\Martabak_Hati_Revisi_Final.docx"
    doc.save(target_path)
    doc.save(clean_path)
    print(f"[SUCCESS] File berhasil diperbaiki dan disimpan di:\n1. {target_path}\n2. {clean_path}")

if __name__ == "__main__":
    fix_document()
